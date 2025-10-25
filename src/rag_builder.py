"""
RAG構築クラス
Databricks 16.4 LTS + 既存ライブラリ完全保護版

RAGシステムの構築・学習に関する具体的な実装を担当
ファイル処理、チャンキング、ベクトル化、インデックス作成などを実装
"""

import os
import json
import pickle
from typing import List, Dict, Any, Tuple
from pathlib import Path

# 標準ライブラリ（Databricks標準搭載）
import pandas as pd
import numpy as np

# 既存ライブラリを使用（バージョン変更禁止）
from langchain_text_splitters import RecursiveCharacterTextSplitter  # 既存: langchain-text-splitters 0.3.8
from langchain_core.documents import Document  # 既存: langchain-core 0.3.51

# 新規インストールライブラリ
import faiss  # faiss-cpu 1.8.0
from rank_bm25 import BM25Okapi  # rank-bm25 0.2.2
import docx  # python-docx 1.1.2
import openpyxl  # openpyxl 3.1.5
from pptx import Presentation  # python-pptx 1.0.2

# 基底クラスをインポート
from base_llm import LLMBase


class RAGBuilder(LLMBase):
    """
    RAG構築クラス
    
    機能：
    - ファイル処理（Word、Excel、PowerPoint、テキスト）
    - チャンキング処理
    - ベクトル化処理
    - Faissインデックス構築
    - BM25インデックス構築
    - Unity Catalogボリュームへの保存
    """
    
    def __init__(self, dbutils, chunk_size: int = 1000, chunk_overlap: int = 200, temperature: float = 0.7):
        """
        RAG構築クラスの初期化
        
        Args:
            dbutils: Databricksのユーティリティオブジェクト
            chunk_size (int): チャンクサイズ（文字数）
            chunk_overlap (int): オーバーラップサイズ
            temperature (float): LLM温度パラメータ
        """
        super().__init__(dbutils, temperature)
        
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        
        # 日本語対応のTextSplitterを設定
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.chunk_size,
            chunk_overlap=self.chunk_overlap,
            separators=["\n\n", "\n", "。", "、", " ", ""]
        )
        
    def process_directory(self, target_directory: str = None) -> Dict[str, Any]:
        """
        指定ディレクトリ内の全ファイルを処理してRAGシステムを構築
        
        Args:
            target_directory (str, optional): 処理対象ディレクトリ（未指定時はINPUT_DIRECTORY）
            
        Returns:
            Dict[str, Any]: 処理結果の統計情報
        """
        if target_directory is None:
            target_directory = self.INPUT_DIRECTORY
            
        print(f"ディレクトリ処理開始: {target_directory}")
        
        # 1. ファイル抽出処理
        extraction_stats = self._extract_texts_from_directory(target_directory)
        
        # 2. チャンキング処理
        chunking_stats = self._chunk_extracted_texts()
        
        # 3. インデックス構築処理
        indexing_stats = self._build_indexes()
        
        # 4. 統計情報の集約
        total_stats = {
            "extraction": extraction_stats,
            "chunking": chunking_stats,
            "indexing": indexing_stats,
            "timestamp": pd.Timestamp.now().isoformat()
        }
        
        # 5. 処理結果を保存
        self.save_config(total_stats, os.path.join(self.VOLUME_PATH, "build_stats.json"))
        
        print("RAG構築処理が完了しました")
        return total_stats
        
    def _extract_texts_from_directory(self, directory: str) -> Dict[str, Any]:
        """
        ディレクトリ内の全ファイルからテキストを抽出
        
        Args:
            directory (str): 処理対象ディレクトリ
            
        Returns:
            Dict[str, Any]: 抽出処理の統計情報
        """
        extracted_files = []
        error_files = []
        
        # サポートされるファイル拡張子
        supported_extensions = {'.docx', '.xlsx', '.pptx', '.txt'}
        
        # ディレクトリを再帰的に探索
        for root, dirs, files in os.walk(directory):
            for file in files:
                file_path = os.path.join(root, file)
                file_extension = Path(file).suffix.lower()
                
                if file_extension in supported_extensions:
                    try:
                        # ファイル形式に応じてテキスト抽出
                        if file_extension == '.docx':
                            text = self._extract_text_from_docx(file_path)
                        elif file_extension == '.xlsx':
                            text = self._extract_text_from_xlsx(file_path)
                        elif file_extension == '.pptx':
                            text = self._extract_text_from_pptx(file_path)
                        elif file_extension == '.txt':
                            text = self._extract_text_from_txt(file_path)
                        else:
                            continue
                            
                        if text.strip():  # 空でないテキストのみ処理
                            # 抽出したテキストを保存
                            relative_path = os.path.relpath(file_path, directory)
                            output_filename = f"{Path(file).stem}.txt"
                            output_path = os.path.join(self.EXTRACTED_TEXT_DIR, output_filename)
                            
                            with open(output_path, 'w', encoding='utf-8') as f:
                                f.write(text)
                                
                            extracted_files.append({
                                "original_file": relative_path,
                                "extracted_file": output_filename,
                                "file_type": file_extension,
                                "text_length": len(text)
                            })
                            
                    except Exception as e:
                        error_files.append({
                            "file": file_path,
                            "error": str(e)
                        })
                        print(f"ファイル処理エラー {file_path}: {str(e)}")
                        
        # メタデータを保存
        metadata = {
            "extracted_files": extracted_files,
            "error_files": error_files,
            "total_extracted": len(extracted_files),
            "total_errors": len(error_files),
            "timestamp": pd.Timestamp.now().isoformat()
        }
        
        self.save_config(metadata, os.path.join(self.EXTRACTED_TEXT_DIR, "metadata.json"))
        
        print(f"テキスト抽出完了: {len(extracted_files)}件成功、{len(error_files)}件エラー")
        return metadata
        
    def _extract_text_from_docx(self, file_path: str) -> str:
        """
        Wordファイルからテキストを抽出
        
        Args:
            file_path (str): Wordファイルのパス
            
        Returns:
            str: 抽出されたテキスト
        """
        try:
            doc = docx.Document(file_path)
            text_parts = []
            
            # 段落からテキストを抽出
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
                    
            # テーブルからテキストを抽出
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_parts.append(cell.text)
                            
            return "\n".join(text_parts)
            
        except Exception as e:
            raise Exception(f"Wordファイル読み込みエラー: {str(e)}")
            
    def _extract_text_from_xlsx(self, file_path: str) -> str:
        """
        Excelファイルからテキストを抽出（全シート対応）
        
        Args:
            file_path (str): Excelファイルのパス
            
        Returns:
            str: 抽出されたテキスト
        """
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            text_parts = []
            
            # 全シートを処理
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                text_parts.append(f"=== シート: {sheet_name} ===")
                
                # セルの内容を抽出
                for row in sheet.iter_rows():
                    row_texts = []
                    for cell in row:
                        if cell.value is not None:
                            row_texts.append(str(cell.value))
                    if row_texts:
                        text_parts.append("\t".join(row_texts))
                        
            return "\n".join(text_parts)
            
        except Exception as e:
            raise Exception(f"Excelファイル読み込みエラー: {str(e)}")
            
    def _extract_text_from_pptx(self, file_path: str) -> str:
        """
        PowerPointファイルからテキストを抽出（スライド + ノート）
        
        Args:
            file_path (str): PowerPointファイルのパス
            
        Returns:
            str: 抽出されたテキスト
        """
        try:
            presentation = Presentation(file_path)
            text_parts = []
            
            for i, slide in enumerate(presentation.slides, 1):
                text_parts.append(f"=== スライド {i} ===")
                
                # スライド内容を抽出
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text_parts.append(shape.text)
                        
                # ノート内容を抽出
                if slide.notes_slide and slide.notes_slide.notes_text_frame:
                    notes_text = slide.notes_slide.notes_text_frame.text.strip()
                    if notes_text:
                        text_parts.append(f"ノート: {notes_text}")
                        
            return "\n".join(text_parts)
            
        except Exception as e:
            raise Exception(f"PowerPointファイル読み込みエラー: {str(e)}")
            
    def _extract_text_from_txt(self, file_path: str) -> str:
        """
        テキストファイルからテキストを抽出
        
        Args:
            file_path (str): テキストファイルのパス
            
        Returns:
            str: 抽出されたテキスト
        """
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except UnicodeDecodeError:
            # UTF-8で読めない場合はshift_jisで試行
            try:
                with open(file_path, 'r', encoding='shift_jis') as f:
                    return f.read()
            except Exception as e:
                raise Exception(f"テキストファイル読み込みエラー: {str(e)}")
                
    def _chunk_extracted_texts(self) -> Dict[str, Any]:
        """
        抽出されたテキストをチャンキング処理
        
        Returns:
            Dict[str, Any]: チャンキング処理の統計情報
        """
        chunked_documents = []
        chunk_metadata = []
        
        # 抽出されたテキストファイルを処理
        for file_name in os.listdir(self.EXTRACTED_TEXT_DIR):
            if file_name.endswith('.txt') and file_name != 'metadata.json':
                file_path = os.path.join(self.EXTRACTED_TEXT_DIR, file_name)
                
                try:
                    with open(file_path, 'r', encoding='utf-8') as f:
                        text = f.read()
                        
                    # LangChainのTextSplitterでチャンキング
                    chunks = self.text_splitter.split_text(text)
                    
                    # Documentオブジェクトとして構造化
                    for i, chunk in enumerate(chunks):
                        doc = Document(
                            page_content=chunk,
                            metadata={
                                "source_file": file_name,
                                "chunk_id": i,
                                "chunk_size": len(chunk)
                            }
                        )
                        chunked_documents.append(doc)
                        
                    chunk_metadata.append({
                        "source_file": file_name,
                        "total_chunks": len(chunks),
                        "total_length": len(text)
                    })
                    
                except Exception as e:
                    print(f"チャンキングエラー {file_name}: {str(e)}")
                    
        # チャンク化されたドキュメントを保存
        chunks_data = []
        for doc in chunked_documents:
            chunks_data.append({
                "content": doc.page_content,
                "metadata": doc.metadata
            })
            
        output_path = os.path.join(self.CHUNKED_TEXT_DIR, "chunks_data.json")
        self.save_config(chunks_data, output_path)
        
        # メタデータを保存
        metadata = {
            "chunk_files": chunk_metadata,
            "total_chunks": len(chunked_documents),
            "chunk_size": self.chunk_size,
            "chunk_overlap": self.chunk_overlap,
            "timestamp": pd.Timestamp.now().isoformat()
        }
        
        self.save_config(metadata, os.path.join(self.CHUNKED_TEXT_DIR, "chunk_metadata.json"))
        
        print(f"チャンキング完了: {len(chunked_documents)}個のチャンクを生成")
        return metadata
        
    def _build_indexes(self) -> Dict[str, Any]:
        """
        FaissとBM25のインデックスを構築
        
        Returns:
            Dict[str, Any]: インデックス構築の統計情報
        """
        # チャンクデータを読み込み
        chunks_file = os.path.join(self.CHUNKED_TEXT_DIR, "chunks_data.json")
        chunks_data = self.load_config(chunks_file)
        
        if not chunks_data:
            raise Exception("チャンクデータが見つかりません")
            
        # テキストとメタデータを分離
        texts = [chunk["content"] for chunk in chunks_data]
        metadatas = [chunk["metadata"] for chunk in chunks_data]
        
        # 1. ベクトル埋め込みを生成
        print("ベクトル埋め込みを生成中...")
        embeddings = self.get_embeddings(texts)
        
        # リスト形式の埋め込みベクトルをNumPy配列に変換
        try:
            embeddings_array = np.array(embeddings, dtype=np.float32)
            # 埋め込みベクトルが1次元になってしまった場合（ドキュメント数が1件）は2次元化
            if embeddings_array.ndim == 1:
                embeddings_array = embeddings_array.reshape(1, -1)
            print(f"埋め込みベクトル形状: {embeddings_array.shape}")
        except Exception as e:
            print(f"埋め込みベクトル変換エラー: {str(e)}")
            raise Exception(f"埋め込みベクトルの変換に失敗しました: {str(e)}")
        
        # 2. Faissインデックスを構築
        print("Faissインデックスを構築中...")
        dimension = embeddings_array.shape[1]
        faiss_index = faiss.IndexFlatIP(dimension)  # 内積類似度
        faiss_index.add(embeddings_array)
        
        # 3. BM25インデックスを構築
        print("BM25インデックスを構築中...")
        tokenized_texts = [text.split() for text in texts]
        bm25_index = BM25Okapi(tokenized_texts)
        
        # 4. インデックスを保存
        faiss_path = os.path.join(self.FAISS_INDEX_DIR, "faiss.index")
        faiss.write_index(faiss_index, faiss_path)
        
        bm25_path = os.path.join(self.FAISS_INDEX_DIR, "bm25_index.pkl")
        with open(bm25_path, 'wb') as f:
            pickle.dump(bm25_index, f)
            
        # 5. ドキュメントとメタデータを保存
        documents_data = {
            "texts": texts,
            "metadatas": metadatas
        }
        documents_path = os.path.join(self.FAISS_INDEX_DIR, "documents.json")
        self.save_config(documents_data, documents_path)
        
        # 6. インデックス設定情報を保存
        index_config = {
            "faiss_dimension": dimension,
            "total_documents": len(texts),
            "chunk_size": self.chunk_size,
            "chunk_overlap": self.chunk_overlap,
            "embedding_model": "bge_m3",
            "timestamp": pd.Timestamp.now().isoformat()
        }
        config_path = os.path.join(self.FAISS_INDEX_DIR, "index_config.json")
        self.save_config(index_config, config_path)
        
        print(f"インデックス構築完了: {len(texts)}個のドキュメント")
        return index_config